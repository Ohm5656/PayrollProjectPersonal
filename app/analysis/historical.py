"""Phase 1: inspect every workbook/sheet and document historical evidence.

Deliberately independent of the Phase 2 attendance parser so its raw historical
observations can serve as an independent regression oracle.
"""
from collections import Counter, defaultdict
from datetime import date, datetime, time
from decimal import Decimal
from difflib import SequenceMatcher
from pathlib import Path
import argparse
import csv
import json
import re
import unicodedata
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from app.excel.workbook_source import (
    WorkbookSnapshot, discover_workbooks, file_hash, filename_month, safe_output_path,
)

LEAVE = re.compile(r'(?<![ก-๙])ลา(?=\s|$|\d|ป่วย|กิจ|พักร้อน|ไม่รับ|ครึ่ง)')
LATE = re.compile(r'(?:เช็[คก]อิน.*(?:เกิน|8:36)|เข้างานสาย|ตอกบัตรเกิน|ออกก่อนเวลา)')
FIELD_LABELS = {'bonus': 'เบี้ยขยัน', 'travel': 'ค่าเดินทาง', 'phone': 'ค่าโทรศัพท์',
                'allowance': 'เบี้ยเลี้ยง', 'position': 'ค่าตำแหน่ง', 'housing': 'ที่พัก+ค่าน้ำ+ไฟ',
                'social_security': 'ประกันสังคม', 'advance': 'เบิกล่วงหน้า',
                'tool_loss': 'หักเครื่องมือหาย', 'loan': 'หักเงินกู้ยืม', 'tax': 'ภาษี ภงด.1',
                'special_bonus': 'โบนัส', 'annual_leave_return': 'คืนพักร้อน',
                'no_leave_award': 'รางวัลไม่ลา', 'advance_return': 'คืนเบิก advance'}


def normalize_name(value):
    value = unicodedata.normalize('NFC', str(value or '')).strip()
    value = re.sub(r'^(?:นาย|นางสาว|น\.ส\.|นาง)\s*', '', value)
    return ' '.join(value.split())


def numeric(value):
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        result = Decimal(str(value))
        return result if result.is_finite() else None
    return None


def holiday_type(text):
    if 'น้ำท่วม' in text:
        return 'EMERGENCY_CLOSURE'
    if 'หยุดแทน' in text:
        return 'SUBSTITUTE_HOLIDAY'
    if 'สงกรานต์' in text:
        return 'COMPANY_SPECIAL_HOLIDAY'
    if 'วันแม่' in text or 'ปีใหม่' in text:
        return 'PUBLIC_HOLIDAY'
    if 'วันอาทิตย์' in text:
        return 'WEEKLY_HOLIDAY'
    return 'UNKNOWN'


def _sheet_info(sheet, cached):
    formulas = []
    errors = []
    for row in sheet:
        for cell in row:
            if cell.data_type == 'f':
                formulas.append({'cell': cell.coordinate, 'formula': cell.value,
                                 'cached': cached[cell.coordinate].value})
            elif cell.data_type == 'e':
                errors.append({'cell': cell.coordinate, 'error': cell.value})
    return {'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column,
            'state': sheet.sheet_state, 'merged_ranges': sorted(map(str, sheet.merged_cells.ranges)),
            'formulas': formulas, 'errors': errors,
            'print_area': str(sheet.print_area), 'print_title_rows': sheet.print_title_rows,
            'orientation': sheet.page_setup.orientation, 'paper_size': sheet.page_setup.paperSize}


def _attendance(sheet, cached, month):
    full_name = sheet['B1'].value
    headers = {c.coordinate: str(c.value) for c in sheet[2] if c.value is not None}
    records = []
    other_periods = Counter()
    for row in sheet.iter_rows(min_row=3):
        value = row[0].value
        if not isinstance(value, (date, datetime)):
            continue
        day = value.date() if isinstance(value, datetime) else value
        if day.month != month or day.year != 2026:
            other_periods[f'{day.year:04}-{day.month:02}'] += 1
            continue
        text_cells = {c.coordinate: c.value for c in row[3:] if isinstance(c.value, str)
                      and c.data_type != 'f'}
        raw = {key: sheet[f'{col}{row[0].row}'].value
               for key, col in [('ot', 'H'), ('allowance', 'I'), ('travel', 'J')]}
        records.append({'date': day.isoformat(), 'row': row[0].row, 'raw': raw,
                        'text_cells': text_cells,
                        'leave': {k: v for k, v in text_cells.items() if LEAVE.search(v)},
                        'late': {k: v for k, v in text_cells.items() if LATE.search(v)},
                        'holiday': {k: v for k, v in text_cells.items() if 'หยุด' in v},
                        'bonus_cut': {k: v for k, v in text_cells.items() if 'ตัดเบี้ยขยัน' in v}})
    sums = {k: sum((numeric(r['raw'][k]) or Decimal(0) for r in records), Decimal(0))
            for k in ('ot', 'allowance', 'travel')}
    totals = []
    for row in sheet:
        if isinstance(row[0].value, (date, datetime)):
            continue
        for c in row:
            if c.column in (8, 9, 10) and c.data_type == 'f' and 'SUM(' in c.value.upper():
                totals.append({'cell': c.coordinate, 'formula': c.value,
                               'cached': cached[c.coordinate].value})
    metadata = sheet['J1'].value
    return {'sheet': sheet.title, 'full_name': full_name,
            'normalized_name': normalize_name(full_name), 'headers': headers,
            'metadata_period_raw': str(metadata), 'records': records, 'numeric_sums': sums,
            'other_periods': dict(other_periods), 'totals': totals,
            'numeric_text': [{'cell': f'{col}{r["row"]}', 'text': r['raw'][key]}
                             for r in records for key, col in [('ot', 'H'), ('allowance', 'I'), ('travel', 'J')]
                             if isinstance(r['raw'][key], str)]}


def _payroll(sheet, cached):
    expected = {'B5': 'ชื่อ-สกุล', 'D5': 'ค่าจ้าง', 'F6': 'ค่าล่วงเวลาปกติ x 1.5',
                'I6': 'ค่าทำงานในวันหยุด', 'L6': 'OT  วันหยุด x 3', 'O7': 'เบี้ยขยัน',
                'P7': 'ค่าเดินทาง', 'R7': 'เบี้ยเลี้ยง'}
    mismatches = [f'{k}: expected {v!r}, got {sheet[k].value!r}'
                  for k, v in expected.items() if sheet[k].value != v]
    fields_by_column = {'base': 4, 'normal_ot': 7, 'holiday_days': 10, 'holiday_ot': 13}
    for key, label in FIELD_LABELS.items():
        found = [c.column for c in sheet[7] if c.value == label]
        if len(found) > 1:
            mismatches.append(f'Ambiguous header: {label}')
        fields_by_column[key] = found[0] if len(found) == 1 else None
    net_headers = [c.coordinate for c in sheet[5] if c.value == 'รายได้สุทธิ']
    if len(net_headers) != 1:
        mismatches.append('Missing/ambiguous net income header')
    rows = []
    for row in sheet:
        number, name = row[0].value, row[1].value
        if not isinstance(number, int) or isinstance(number, bool) or not isinstance(name, str):
            continue
        fields = {}
        for key, col in fields_by_column.items():
            if col is None:
                fields[key] = {'cell': None, 'raw': None, 'value': None, 'formula': None, 'present': False}
                continue
            cell = sheet.cell(row[0].row, col)
            value = cached[cell.coordinate].value if cell.data_type == 'f' else cell.value
            fields[key] = {'cell': cell.coordinate, 'raw': cell.value, 'value': value,
                           'formula': cell.value if cell.data_type == 'f' else None, 'present': True}
        rows.append({'row': row[0].row, 'name': name, 'normalized_name': normalize_name(name),
                     'fields': fields})
    return {'sheet': sheet.title, 'period_text': sheet['A3'].value,
            'template_issues': mismatches, 'net_income_header': net_headers,
            'field_columns': fields_by_column,
            'headers': {c.coordinate: c.value for row in sheet.iter_rows(min_row=5, max_row=7)
                        for c in row if c.value is not None}, 'employees': rows}


def inspect_all(reference_root):
    result = {'files': [], 'comparisons': [], 'unmatched_payroll': [], 'read_only_verified': False}
    sources = [(kind, path) for kind in ('attendance', 'payroll')
               for path in discover_workbooks(reference_root / kind)]
    initial = {str(path): file_hash(path) for _, path in sources}
    for kind, path in sources:
        snapshot = WorkbookSnapshot.read(path)
        month = filename_month(path)
        entry = {'kind': kind, 'file': path.relative_to(reference_root).as_posix(),
                 'month': month, 'sha256': snapshot.sha256, 'bytes': len(snapshot.content)}
        with snapshot.open() as (book, cached):
            entry['sheets'] = [_sheet_info(s, cached[s.title]) for s in book]
            entry['external_links'] = len(book._external_links)
            entry['defined_name_count'] = len(book.defined_names)
            if kind == 'attendance':
                entry['attendance'] = [_attendance(s, cached[s.title], month) for s in book]
            else:
                main = [s for s in book if s.title.strip() == 'คิดค่าจ้าง']
                entry['payroll'] = _payroll(main[0], cached[main[0].title]) if len(main) == 1 else None
                entry['payslips'] = [{'sheet': s.title,
                                     'employee_label': re.split(r'\(|เลขที่|เลขบัญชี', str(s['C3'].value or ''))[0].strip(),
                                     'fields': {c.coordinate: c.value for row in s.iter_rows(min_row=6, max_row=14)
                                                for c in row if c.column in (2, 4, 7, 9) and c.value is not None}}
                                    for s in book if 'PAY SLIP' in str(s['B1'].value)]
        result['files'].append(entry)
    grouped = defaultdict(lambda: defaultdict(list))
    for f in result['files']:
        grouped[f['month']][f['kind']].append(f)
    for month, pair in sorted(grouped.items()):
        if len(pair['attendance']) != 1 or len(pair['payroll']) != 1:
            continue  # Ambiguous pairing is visible in the inventory, never guessed.
        attendance, payroll = pair['attendance'][0], pair['payroll'][0]['payroll']
        if not payroll or payroll['template_issues']:
            continue
        matched = set()
        for sheet in attendance['attendance']:
            name = sheet['normalized_name']
            candidates = [e for e in payroll['employees'] if e['normalized_name'] == name and name]
            duplicate_names = sum(a['normalized_name'] == name for a in attendance['attendance'])
            conflicts = bool(name and sheet['sheet'].strip('- ') and
                             normalize_name(sheet['sheet']) != name.split()[0])
            exact = (len(candidates) == 1 and duplicate_names == 1 and not conflicts
                     and bool(sheet['records']) and not sheet['other_periods'])
            suggestions = sorted(payroll['employees'],
                                 key=lambda e: SequenceMatcher(None, name, e['normalized_name']).ratio(),
                                 reverse=True)[:2] if name else []
            employee = candidates[0] if exact else None
            if exact:
                matched.add(employee['row'])
            comparison = {'month': month, 'sheet': sheet['sheet'], 'attendance_name': sheet['full_name'],
                          'mapping': 'EXACT_FULL_NAME' if exact else 'REVIEW_REQUIRED',
                          'name_conflict': conflicts, 'duplicate_identity': duplicate_names > 1 and bool(name),
                          'payroll_name': employee['name'] if employee else None,
                          'suggestions_only': [e['name'] for e in suggestions],
                          'daily_rows': len(sheet['records']),
                          'leave_dates': [r['date'] for r in sheet['records'] if r['leave']],
                          'late_dates': [r['date'] for r in sheet['records'] if r['late']],
                          'bonus_cut_dates': [r['date'] for r in sheet['records'] if r['bonus_cut']]}
            for key in ('travel', 'allowance'):
                source = sheet['numeric_sums'][key]
                observed = employee['fields'][key]['value'] if employee else None
                payroll_number = Decimal(0) if employee and observed is None else numeric(observed)
                comparison[f'attendance_{key}'] = source
                comparison[f'payroll_{key}'] = payroll_number
                comparison[f'{key}_difference'] = source - payroll_number if payroll_number is not None else None
            comparison['recorded_bonus'] = employee['fields']['bonus']['value'] if employee else None
            differences = [comparison[f'{k}_difference'] for k in ('travel', 'allowance')]
            comparison['difference_classification'] = ('unresolved identity' if not exact else
                'missing source data' if not sheet['records'] else
                'unknown — confirm exception/manual override' if any(d not in (None, 0) for d in differences) else
                'numeric agreement; text notes still require review' if sheet['numeric_text'] else 'numeric agreement')
            result['comparisons'].append(comparison)
        result['unmatched_payroll'].extend({'month': month, 'name': e['name']}
                                           for e in payroll['employees'] if e['row'] not in matched)
    if any(file_hash(Path(p)) != h for p, h in initial.items()):
        raise RuntimeError('Reference content changed during analysis; report is not a stable baseline')
    final_paths = {(kind, str(p)) for kind in ('attendance', 'payroll')
                   for p in discover_workbooks(reference_root / kind)}
    if final_paths != {(kind, str(p)) for kind, p in sources}:
        raise RuntimeError('Reference inventory changed during analysis')
    result['read_only_verified'] = True
    return result


def _json_default(value):
    if isinstance(value, ArrayFormula):
        return {'type': 'array_formula', 'ref': value.ref, 'text': value.text}
    if isinstance(value, DataTableFormula):
        return {'type': 'data_table_formula', **dict(value)}
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    raise TypeError(type(value).__name__)


def _table(headers, rows):
    def escape(value):
        return str(value if value is not None else '—').replace('|', '\\|').replace('\n', ' ')
    return ['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join('---' for _ in headers) + ' |'] + [
        '| ' + ' | '.join(escape(v) for v in row) + ' |' for row in rows]


def render_report(data):
    lines = ['# NTP historical payroll analysis', '',
             'Scope: all locally supplied `.xlsx` files in `reference_data/attendance` and `reference_data/payroll`.',
             'Analysis baseline: 11 October 2026 (Asia/Bangkok). No Excel COM, workbook saves, or Excel processes are used.',
             'Every worksheet and every populated formula cell was inspected from a saved byte snapshot. '
             'SHA-256 was verified before and after inspection; all source content was unchanged.',
             'Daily records use Gregorian 2026 dates. The literal 2569 year in J1 is Buddhist-era metadata; '
             'it is not an instruction to interpret daily dates as year 2569.', '',
             'The report contains confidential employee names and payroll values. Keep it local. '
             'National IDs and bank account details are excluded.', '',
             '## Inventory and worksheet structures', '']
    lines += _table(['Kind', 'Month', 'File', 'Sheets', 'SHA-256'],
                    [(f['kind'], f['month'], f['file'], len(f['sheets']), f['sha256']) for f in data['files']])
    for f in data['files']:
        lines += ['', f'### {f["file"]}', '']
        lines += _table(['Exact sheet name', 'Used size', 'Merges', 'Formulas', 'Cached formula values missing', 'State'],
                        [(f'`{s["name"]}`', f'{s["rows"]} × {s["columns"]}', len(s['merged_ranges']),
                          len(s['formulas']), sum(v['cached'] is None for v in s['formulas']), s['state']) for s in f['sheets']])
        lines += [f'External links: {f["external_links"]}; defined names: {f["defined_name_count"]}. '
                  'Full formula coordinates, cached values, merge ranges and print settings are in `historical_analysis.json`.']
        if f['kind'] == 'attendance':
            lines += [''] + _table(['Sheet', 'B1 employee', '2026 rows in named month', 'Other date periods', 'Numeric H / I / J sums'],
                [(a['sheet'], a['full_name'], len(a['records']), a['other_periods'],
                  ' / '.join(str(a['numeric_sums'][k]) for k in ('ot', 'allowance', 'travel'))) for a in f['attendance']])
            signatures = {tuple(a['headers'].items()) for a in f['attendance']}
            lines += ['Header variants:'] + [f'- {dict(sig)}' for sig in sorted(signatures)]
            totals = {(t['cell'], t['formula']) for a in f['attendance'] for t in a['totals']}
            lines += ['Observed total formulas (row locations vary): ' + '; '.join(f'`{c}: {v}`' for c, v in sorted(totals))]
        elif f['payroll']:
            lines += ['', f'Payroll period label: {f["payroll"]["period_text"]}',
                      f'Template mismatches: {f["payroll"]["template_issues"] or "none for inspected signature"}.',
                      'Multi-row payroll headers: ' + str(f['payroll']['headers']),
                      'Payslip sheets: ' + ', '.join(f'`{p["sheet"]}`' for p in f['payslips']) + '. '
                      'Payslip amounts are frequently literal numbers, so recalculating payroll alone cannot update them.']
    lines += ['', '## Identity evidence and month differences', '',
              'Only unique exact full-name matches (honorifics and repeated whitespace normalized) are used for comparisons. '
              'Sheet/B1 conflicts and duplicate names block matching. Similar names below are suggestions, never aliases.', '']
    lines += _table(['Month', 'Attendance sheet / B1', 'Mapping', 'Payroll name or candidate suggestions'],
                    [(c['month'], f'{c["sheet"]} / {c["attendance_name"]}', c['mapping'],
                      c['payroll_name'] or '; '.join(c['suggestions_only'])) for c in data['comparisons']])
    lines += ['', 'Confirmed structural observations:',
              '- `ภานุวัฒน์` attendance versus payroll `ภาณุวัฒน์ มูลที`: spelling differs and attendance omits the surname. Requires mapping confirmation.',
              '- B1 `อดิศักดิ์` is first-name-only in months 7 and 9. Sheet `--` retains `อดิศักดิ์ สาคำตา` with stale April dates in months 5–9. Neither can automatically establish the current payroll identity.',
              '- Month 8 sheet `อดิศักดิ์` has B1 `บรรจง มูลที`, also present on sheet `บรรจง`. This identity conflict must be resolved before money is transferred.',
              '- Sheet `--` has April data (current in April, stale thereafter); sheet `-` has no employee name and October 2021 dates. Stale sheets are inspected and quarantined, not treated as zero-paid employees.',
              '- January net income is AD, March net income is AB, and the other months use AA. March adds a separate bonus column; January includes leave-return / no-leave-award / advance-return columns. Deductions must be located by headers, not fixed letters.',
              '- A missing/added sheet is an observation, not proof of joining, leaving, or employment status.', '']
    previous_a, previous_p = set(), set()
    changes = []
    for month in sorted({f['month'] for f in data['files']}):
        a = {s['sheet'] for f in data['files'] if f['kind'] == 'attendance' and f['month'] == month for s in f['attendance']}
        p = {e['normalized_name'] for f in data['files'] if f['kind'] == 'payroll' and f['month'] == month and f['payroll'] for e in f['payroll']['employees']}
        changes.append((month, ', '.join(sorted(a - previous_a)), ', '.join(sorted(previous_a - a)),
                        ', '.join(sorted(p - previous_p)), ', '.join(sorted(previous_p - p))))
        previous_a, previous_p = a, p
    lines += _table(['Month', 'Attendance sheets added', 'Attendance sheets removed', 'Payroll names added', 'Payroll names removed'], changes)
    lines += ['', 'Unmatched payroll names (includes directors and spelling/identity review cases):', '']
    lines += _table(['Month', 'Payroll employee'], [(u['month'], u['name']) for u in data['unmatched_payroll']])
    lines += ['', '## Travel and allowance comparisons', '',
              'Difference = sum of numeric daily attendance values minus existing payroll value. '
              'Blank payroll travel/allowance is shown as zero for this historical comparison only. '
              'Text is not converted to money; unresolved identities have no payroll comparison. '
              '2021/stale-period rows and aggregate SUM rows are excluded. No historical figures are changed.', '']
    lines += _table(['Month', 'Employee sheet', 'Travel daily', 'Payroll P', 'Δ travel', 'Allowance daily', 'Payroll R', 'Δ allowance', 'Classification'],
                    [(c['month'], c['sheet'], c['attendance_travel'], c['payroll_travel'], c['travel_difference'],
                      c['attendance_allowance'], c['payroll_allowance'], c['allowance_difference'], c['difference_classification'])
                     for c in data['comparisons']])
    matched = [c for c in data['comparisons'] if c['mapping'] == 'EXACT_FULL_NAME']
    lines += ['', f'Unique exact-name comparisons: {len(matched)}. Travel mismatches: '
              f'{sum(c["travel_difference"] not in (None, 0) for c in matched)}; allowance mismatches: '
              f'{sum(c["allowance_difference"] not in (None, 0) for c in matched)}.',
              'Mismatch causes remain unknown unless the source establishes a missing record. '
              'Manual overrides or valid business exceptions are hypotheses requiring confirmation, not inferred classifications.']
    lines += ['', '## Attendance bonus versus leave and late/check-in', '',
              'Leave detection uses a leave token, not the substring `ลา` inside `เวลา`, `สงขลา` or `ลาดหลุมแก้ว`. '
              'Late/check-in detection requires an explicit violation note; a clock time alone does not establish a violation. '
              'Blank bonus cells mean no bonus recorded, not a confirmed eligibility rule.', '']
    lines += _table(['Month', 'Employee', 'Leave dates', 'Late/check-in dates', 'Explicit bonus-cut dates', 'Recorded O bonus'],
                    [(c['month'], c['sheet'], ', '.join(c['leave_dates']), ', '.join(c['late_dates']),
                      ', '.join(c['bonus_cut_dates']), c['recorded_bonus']) for c in matched
                     if c['leave_dates'] or c['late_dates'] or c['bonus_cut_dates']])
    counterexamples = [c for c in matched if c['late_dates'] and (numeric(c['recorded_bonus']) or 0) > 0]
    leave_cases = [c for c in matched if c['leave_dates']]
    lines += ['', f'Leave employee-months with exact identity: {len(leave_cases)}; '
              f'positive recorded bonus in {sum((numeric(c["recorded_bonus"]) or 0) > 0 for c in leave_cases)}.',
              f'Late/check-in employee-months with a positive bonus: {len(counterexamples)}. '
              'These observations rule out a universal `late => zero attendance bonus` rule. '
              'Allowance-cut notes and attendance-bonus rules must remain separate.',
              'Leave type, paid/unpaid status, hours/days conversion and bonus consequences still require human confirmation.']
    lines += ['', '## Source events and special holidays', '',
              'Holiday types are textual suggestions. A note on one employee sheet does not establish a company-wide closure. '
              'Mixed holiday/work descriptions and conflicting sheets require review. No calendar is confirmed automatically.', '']
    holiday_groups = defaultdict(set)
    event_rows, text_rows = [], []
    for f in data['files']:
        if f['kind'] != 'attendance':
            continue
        for a in f['attendance']:
            text_rows += [(f['month'], a['sheet'], t['cell'], t['text']) for t in a['numeric_text']]
            for r in a['records']:
                for cell, text in r['holiday'].items():
                    holiday_groups[(r['date'], text, holiday_type(text))].add(a['sheet'])
                for category in ('leave', 'late', 'bonus_cut'):
                    event_rows += [(f['month'], a['sheet'], r['date'], cell, category, text)
                                   for cell, text in r[category].items()]
    lines += _table(['Date', 'Original holiday text', 'Suggested type', 'Employee sheets'],
                    [(day, text, kind, ', '.join(sorted(sheets))) for (day, text, kind), sheets in sorted(holiday_groups.items())])
    lines += ['', 'Leave, late and explicit bonus-cut evidence:', '']
    lines += _table(['Month', 'Sheet', 'Date', 'Cell', 'Event', 'Original text'], event_rows)
    lines += ['', 'Non-numeric values in H/I/J (never silently zero for production):', '']
    lines += _table(['Month', 'Sheet', 'Cell', 'Original text'], text_rows)
    lines += ['', '## OT evidence', '',
              'Attendance H totals alone do not establish normal OT, holiday day count or holiday OT. '
              'The table compares the independent historical quantities without reclassifying H.', '']
    ot_rows = []
    for c in matched:
        p = next(f['payroll'] for f in data['files'] if f['kind'] == 'payroll' and f['month'] == c['month'])
        e = next(e for e in p['employees'] if e['name'] == c['payroll_name'])
        a = next(a for f in data['files'] if f['kind'] == 'attendance' and f['month'] == c['month'] for a in f['attendance'] if a['sheet'] == c['sheet'])
        values = [e['fields'][k]['value'] for k in ('normal_ot', 'holiday_days', 'holiday_ot')]
        ot_rows.append((c['month'], c['sheet'], a['numeric_sums']['ot'], *values))
    lines += _table(['Month', 'Employee', 'Numeric H total', 'Payroll G normal OT hours', 'Payroll J holiday days', 'Payroll M holiday OT hours'], ot_rows)
    lines += ['', '## Recurring versus variable payroll evidence', '',
              'These are observations across monthly payroll rows, not approved reuse policies. '
              'Literal and formula values are read with their cached results; missing caches remain unknown. '
              'Base amounts sometimes contain manual additions or daily-rate calculations, so treating D as a universal fixed salary is unsafe.', '']
    histories = defaultdict(list)
    for f in data['files']:
        if f['kind'] == 'payroll' and f['payroll']:
            for e in f['payroll']['employees']:
                for field in ('base', *FIELD_LABELS):
                    item = e['fields'][field]
                    if item['present']:
                        histories[(e['normalized_name'], field)].append((f['month'], item['value']))
    lines += _table(['Employee', 'Field', 'Month:value observations', 'Observed pattern'],
                    [(name, field, '; '.join(f'{m}:{v if v is not None else "blank/uncached"}' for m, v in sorted(values)),
                      'repeated observed value' if len({str(v) for _, v in values}) == 1 else 'variable')
                     for (name, field), values in sorted(histories.items())])
    lines += ['', '## Unknown business rules and recommended boundaries', '',
              '| Field / case | Recommendation |', '| --- | --- |',
              '| Daily numeric travel and allowance | Deterministic sum from validated daily rows; block readiness on text/errors/identity conflicts. Compare stored totals; differences are evidence. |',
              '| Normal OT / holiday work / holiday OT | Keep raw H and work/holiday evidence; human-confirm classification until rules are established. |',
              '| Generic leave and partial leave | Review type, duration, paid status and independent bonus effect. Never infer unpaid leave. |',
              '| Check-in/late / early departure | Preserve explicit notes; review allowance effects separately from bonus eligibility. |',
              '| Company calendar | Suggest textual types; confirm employee scope, required workdays and bonus effects. No external holiday calendar is assumed. |',
              '| Employee spelling and duplicate/conflicting B1 | Require explicit alias mapping; persist stable IDs in Phase 3. No fuzzy auto-match. |',
              '| Salary / phone / position / recurring housing | Show previous observed value; approve employee-specific reuse. Do not overwrite D formulas. |',
              '| Advances / loans / tool losses / bonuses / tax / special adjustments | Require monthly confirmation. Never automatically carry one-time deductions forward. |',
              '| Blank days / stale named sheets / directors without attendance | Review missing data or applicability; absence does not imply zero pay. |',
              '| Payslip literal values and employee-specific rows | Validate label-based adapters and identities before Phase 7 writes; preserve existing sheets/layout. |',
              '| Formula caches | Historical snapshots only; no freshness guarantee or formula evaluator. Excel recalculation is a later isolated COM step. |', '',
              'Phase 1 is complete when this report and its evidence artifacts exist and read-only tests pass. '
              'Phase 2 may then implement parsing only. Database, payroll decisions, desktop UI, writer, payslip updates, '
              'watcher and startup registration are later phases. No COM integration test is appropriate until a writer exists.', '']
    return '\n'.join(lines)


def write_reports(data, output_dir, reference_root):
    outputs = [safe_output_path(Path(output_dir) / filename, reference_root) for filename in
               ('historical_analysis.md', 'historical_analysis.json', 'historical_comparison.csv')]
    # Validate/serialize before touching outputs, including Excel array formulas.
    report_text = render_report(data)
    evidence_text = json.dumps(data, ensure_ascii=False, indent=2, default=_json_default) + '\n'
    outputs[0].parent.mkdir(parents=True, exist_ok=True)
    outputs[0].write_text(report_text, encoding='utf-8')
    outputs[1].write_text(evidence_text, encoding='utf-8')
    with outputs[2].open('w', encoding='utf-8-sig', newline='') as stream:
        keys = list(data['comparisons'][0]) if data['comparisons'] else ['month', 'sheet']
        writer = csv.DictWriter(stream, fieldnames=keys)
        writer.writeheader()
        for comparison in data['comparisons']:
            # Prevent Excel interpreting source-derived strings as CSV formulas.
            writer.writerow({k: "'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v
                             for k, v in comparison.items()})


def main():
    root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description='Read-only historical XLSX analysis; never opens Excel')
    parser.add_argument('--reference', type=Path, default=root / 'reference_data')
    parser.add_argument('--output', type=Path, default=root / 'reports')
    args = parser.parse_args()
    data = inspect_all(args.reference.resolve(strict=True))
    write_reports(data, args.output, args.reference)
    print(f'Inspected {len(data["files"])} workbooks and {sum(len(f["sheets"]) for f in data["files"])} sheets. Source hashes unchanged.')


if __name__ == '__main__':
    main()
