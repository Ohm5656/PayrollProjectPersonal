"""Safe monthly workbook creation from the immutable accounting master.

This module deliberately uses labels, not hard-coded coordinates.  It is the
only workflow component that writes a newly created monthly workbook, and the
development safety boundary confines that copy to ``test_workspace``.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook

from app.excel.payroll_reader import label
from app.safety import MASTER_TEMPLATE_RELATIVE, REFERENCE_ROOT, TEST_ROOT, SafetyError, development_target


THAI_MONTHS = (
    'มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน',
    'กรกฎาคม', 'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม',
)


class MonthlyWorkflowError(ValueError):
    pass


@dataclass(frozen=True)
class MonthlyPayroll:
    year: int
    month: int
    path: Path
    created: bool

    @property
    def payment_date(self) -> date:
        return date(self.year + (self.month == 12), (self.month % 12) + 1, 5)

    @property
    def period_caption(self) -> str:
        paid = self.payment_date
        return (f'ประจำเดือน {THAI_MONTHS[self.month - 1]}  {self.year + 543}  '
                f'จ่ายวันที่ {paid.day} {THAI_MONTHS[paid.month - 1]} {paid.year + 543}')


def master_template_path() -> Path:
    path = (REFERENCE_ROOT / MASTER_TEMPLATE_RELATIVE).resolve()
    if not path.is_file():
        raise MonthlyWorkflowError('ไม่พบแม่แบบรายเดือนที่กำหนด')
    return path


def _inside(path: Path, folder: Path) -> bool:
    resolved, base = path.resolve(), folder.resolve()
    return resolved != base and base in resolved.parents


def validate_workflow_configuration(attendance_folder, payroll_folder, template_path):
    """Validate persisted paths without creating, editing, or opening Excel."""
    attendance, payroll, template = Path(attendance_folder).resolve(), Path(payroll_folder).resolve(), Path(template_path).resolve()
    if not attendance.is_dir() or not payroll.is_dir():
        raise MonthlyWorkflowError('โฟลเดอร์บันทึกเวลาและโฟลเดอร์เงินเดือนต้องมีอยู่แล้ว')
    # Production mode remains disabled.  Folder monitoring/output therefore
    # remains constrained to test_workspace, while the immutable reference
    # master is an allowed read-only input.
    if not _inside(attendance, TEST_ROOT) and attendance != TEST_ROOT.resolve():
        raise SafetyError('Development attendance folder must be inside test_workspace')
    if not _inside(payroll, TEST_ROOT) and payroll != TEST_ROOT.resolve():
        raise SafetyError('Development payroll folder must be inside test_workspace')
    if template != master_template_path() and not _inside(template, TEST_ROOT):
        raise SafetyError('Template must be the read-only master or a test_workspace copy')
    if template.suffix.lower() != '.xlsx' or not template.is_file():
        raise MonthlyWorkflowError('แม่แบบต้องเป็นไฟล์ .xlsx ที่อ่านได้')
    return attendance, payroll, template


def _formulas(book):
    def signature(value):
        # openpyxl represents array formulas as distinct objects after each
        # load, so compare their stored reference/text rather than identity.
        return (getattr(value, 'ref', None), getattr(value, 'text', value))
    return {(sheet.title, cell.coordinate): signature(cell.value)
            for sheet in book.worksheets for row in sheet.iter_rows() for cell in row
            if cell.data_type == 'f'}


def _layout(book):
    return {sheet.title: {
        'merged': tuple(sorted(str(rng) for rng in sheet.merged_cells.ranges)),
        'print_area': str(sheet.print_area),
        'orientation': sheet.page_setup.orientation,
        'paper': sheet.page_setup.paperSize,
    } for sheet in book.worksheets}


def _one(cells, description):
    if len(cells) != 1:
        raise MonthlyWorkflowError(f'พบ {description} {len(cells)} ตำแหน่ง; ต้องพบเพียงตำแหน่งเดียว')
    return cells[0]


def _summary_header_cell(book):
    candidates = []
    for sheet in book.worksheets:
        if label(sheet.title) != label('คิดค่าจ้าง'):
            continue
        names = [cell for row in sheet.iter_rows() for cell in row if label(cell.value) == label('ชื่อ-สกุล')]
        nets = [cell for row in sheet.iter_rows() for cell in row if label(cell.value) == label('รายได้สุทธิ')]
        period = [cell for row in sheet.iter_rows() for cell in row
                  if isinstance(cell.value, str) and 'ประจำเดือน' in cell.value]
        if len(names) == len(nets) == len(period) == 1:
            candidates.append(period[0])
    return _one(candidates, 'หัวกระดาษสรุปเงินเดือน')


def _right_input(sheet, marker, description):
    labels = [cell for row in sheet.iter_rows() for cell in row if label(cell.value) == label(marker)]
    source = _one(labels, description)
    values = [sheet.cell(source.row, col) for col in range(source.column + 1, min(source.column + 8, sheet.max_column + 1))
              if sheet.cell(source.row, col).value is not None]
    return _one(values, description + ' value')


def _payslip_metadata_cells(book):
    result = []
    for sheet in book.worksheets:
        marker = [cell for row in sheet.iter_rows() for cell in row if 'pay slip' in label(cell.value)]
        if len(marker) == 1:
            result.append((sheet, _right_input(sheet, 'ประจำงวด (PERIOD)', 'PERIOD'),
                           _right_input(sheet, 'วันที่จ่าย (DATE)', 'DATE')))
        elif marker:
            raise MonthlyWorkflowError(f'แผ่นงาน {sheet.title} มี PAY SLIP ซ้ำ')
    if not result:
        raise MonthlyWorkflowError('ไม่พบใบเงินเดือนในแม่แบบ')
    return result


def verify_monthly_copy(path, year, month):
    """Return validated copy metadata; reject stale or ambiguous monthly fields."""
    path = Path(path)
    expected = MonthlyPayroll(year, month, path, False)
    book = load_workbook(path, data_only=False, read_only=False)
    try:
        header = _summary_header_cell(book)
        if header.value != expected.period_caption:
            raise MonthlyWorkflowError('หัวกระดาษเงินเดือนยังเป็นงวดเดิม')
        metadata = _payslip_metadata_cells(book)
        for sheet, period_cell, date_cell in metadata:
            if period_cell.value != month:
                raise MonthlyWorkflowError(f'ใบเงินเดือน {sheet.title} มี PERIOD ไม่ตรงงวด')
            value = date_cell.value.date() if isinstance(date_cell.value, datetime) else date_cell.value
            if value != expected.payment_date:
                raise MonthlyWorkflowError(f'ใบเงินเดือน {sheet.title} มี DATE ไม่ตรงวันจ่าย')
        return {'path': str(path), 'header': header.coordinate, 'payslips': len(metadata),
                'payment_date': expected.payment_date.isoformat(), 'period_caption': expected.period_caption}
    finally:
        book.close()


def create_monthly_payroll(template_path, payroll_folder, year, month) -> MonthlyPayroll:
    """Copy the master then replace only deterministic month metadata.

    Existing month files are never overwritten.  Formula text, merged ranges,
    and print setup are compared against the master after the metadata update.
    """
    if not isinstance(year, int) or not isinstance(month, int) or not 1 <= month <= 12:
        raise MonthlyWorkflowError('เดือน/ปีเงินเดือนไม่ถูกต้อง')
    _, folder, template = validate_workflow_configuration(TEST_ROOT, payroll_folder, template_path)
    destination = development_target(folder / f'เดือน {month}.xlsx')
    if destination.exists():
        verify_monthly_copy(destination, year, month)
        return MonthlyPayroll(year, month, destination, False)
    # This is a byte-for-byte new copy before openpyxl changes its three
    # deterministic metadata fields.  The read-only source remains untouched.
    copy2(template, destination)
    source = load_workbook(template, data_only=False, read_only=False)
    try:
        formulae, layout = _formulas(source), _layout(source)
    finally:
        source.close()
    item = MonthlyPayroll(year, month, destination, True)
    book = load_workbook(destination, data_only=False, read_only=False)
    try:
        _summary_header_cell(book).value = item.period_caption
        for _, period_cell, date_cell in _payslip_metadata_cells(book):
            period_cell.value = month
            date_cell.value = item.payment_date
        book.save(destination)
    finally:
        book.close()
    result = load_workbook(destination, data_only=False, read_only=False)
    try:
        if _formulas(result) != formulae or _layout(result) != layout:
            raise MonthlyWorkflowError('การสร้างสำเนาเปลี่ยนสูตร รูปแบบพิมพ์ หรือเซลล์รวม')
    finally:
        result.close()
    verify_monthly_copy(destination, year, month)
    return item
