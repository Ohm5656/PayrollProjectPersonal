from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.db.database import Database
from app.db.repositories import Repository
from app.domain import PayrollState
from app.services.payroll_service import PayrollService, PayrollStateError
from tests.phase_e_fixtures import create_synthetic_excel_fixture, uat_root


def _service_with_configured_synthetic_inputs(label):
    fixture = create_synthetic_excel_fixture(uat_root(label))
    repo = Repository(Database(fixture["root"] / "data" / "uat.db"))
    service = PayrollService(repo)
    initial = service.prepare(fixture["attendance"], fixture["payroll"], year=2026, month=9)
    for person in repo.employees():
        repo.confirm_profile(
            person["id"], base_salary="30000", attendance_bonus_default="1000",
            phone_allowance="500", position_allowance="700", housing="900",
            attendance_required=1, actor="phase-e-uAT",
        )
    for day in range(1, 31):
        repo.confirm_calendar(
            date(2026, 9, day).isoformat(), "WORKING_DAY", "Synthetic UAT calendar",
            required_workday=True, actor="phase-e-uAT", source="synthetic-uAT",
        )
    return fixture, repo, service, initial


def test_phase_e_blocks_unresolved_identity_then_confirms_only_ready_synthetic_draft():
    fixture, repo, service, initial = _service_with_configured_synthetic_inputs("review_and_confirm")
    assert initial["draft"].state == PayrollState.NEEDS_REVIEW
    assert any(issue.kind == "FIELD" for issue in initial["draft"].issues)
    with pytest.raises(PayrollStateError):
        service.confirm(2026, 9)

    prepared = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)
    draft = prepared["draft"]
    assert draft.state == PayrollState.READY
    first = draft.employees[0].values
    assert first["travel"].final_value == Decimal("300")
    assert first["allowance"].final_value == Decimal("600")
    assert first["bonus"].final_value == Decimal("1000")
    assert first["normal_ot"].final_value == Decimal("0")
    assert first["income"].final_value == Decimal("33100")
    assert first["net"].final_value == Decimal("32200")
    assert first["income"].reason.startswith("Preview of validated existing Excel arithmetic")

    confirmed = service.confirm(2026, 9, actor="phase-e-uAT", expected_revision=prepared["run"]["revision"])
    assert confirmed["state"] == PayrollState.CONFIRMED
    assert repo.audit_history(limit=1)[0]["action"] == "PAYROLL_CONFIRMED"
    assert not repo.setting("monitoring_enabled", False)


def test_phase_e_ot_requires_explicit_decision_before_confirmation():
    fixture, repo, service, _ = _service_with_configured_synthetic_inputs("ot_gate")
    book = load_workbook(fixture["attendance"])
    try:
        book["Arun"]["G4"] = 2
        book.save(fixture["attendance"])
    finally:
        book.close()

    blocked = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)
    issue = next(issue for issue in blocked["draft"].issues if issue.kind == "OT")
    assert blocked["draft"].state == PayrollState.NEEDS_REVIEW
    with pytest.raises(PayrollStateError):
        service.confirm(2026, 9)

    repo.decide(blocked["draft"].source_hash, issue.key, {"category": "NORMAL", "hours": "2"},
                reason="Synthetic UAT explicit OT classification", actor="phase-e-uAT")
    ready = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)
    arun = next(employee for employee in ready["draft"].employees if employee.name == "Arun Demo")
    assert ready["draft"].state == PayrollState.READY
    assert arun.values["normal_ot"].final_value == Decimal("2")
    assert arun.values["holiday_ot"].final_value == Decimal("0")


def test_phase_e_unmatched_attendance_identity_stays_blocked():
    fixture = create_synthetic_excel_fixture(uat_root("identity_gate"))
    book = load_workbook(fixture["attendance"])
    try:
        book["Arun"]["B1"] = "Unknown Synthetic Person"
        book.save(fixture["attendance"])
    finally:
        book.close()
    repo = Repository(Database(fixture["root"] / "data" / "uat.db"))
    result = PayrollService(repo).prepare(fixture["attendance"], fixture["payroll"], year=2026, month=9)
    assert result["draft"].state == PayrollState.NEEDS_REVIEW
    assert any(issue.kind == "IDENTITY" for issue in result["draft"].issues)
