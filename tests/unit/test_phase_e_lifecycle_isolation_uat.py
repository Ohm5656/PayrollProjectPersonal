from pathlib import Path

import pytest

from app.domain import PayrollState
from app.excel.attendance_reader import read_attendance
from app.excel.payroll_reader import PayrollWorkbookAdapter
from app.excel.workbook_source import WorkbookReadError
from app.services.payroll_service import PayrollStateError
from tests.phase_e_fixtures import create_synthetic_excel_fixture, uat_root
from tests.unit.test_phase_e_calculation_uat import _service_with_configured_synthetic_inputs


def test_phase_e_multiple_period_fixture_and_invalid_workbook_stay_isolated():
    october = create_synthetic_excel_fixture(uat_root("october"), year=2026, month=10)
    attendance = read_attendance(october["attendance"], year=2026, month=10)
    payroll = PayrollWorkbookAdapter.read(october["payroll"])
    assert (attendance.year, attendance.month) == payroll.period() == (2026, 10)
    assert len(attendance.employees[0].days) == 31
    assert all(path.resolve().is_relative_to(october["root"]) for path in october.values() if isinstance(path, Path))

    invalid = october["root"] / "attendance" / "invalid_10.xlsx"
    invalid.write_text("not an xlsx package", encoding="utf-8")
    with pytest.raises(WorkbookReadError):
        read_attendance(invalid, year=2026, month=10)
    with pytest.raises(WorkbookReadError):
        PayrollWorkbookAdapter.read(invalid)


def test_phase_e_confirmed_source_change_is_protected_and_audited():
    fixture, repo, service, _ = _service_with_configured_synthetic_inputs("source_change")
    prepared = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)
    service.confirm(2026, 9, actor="phase-e-uAT", expected_revision=prepared["run"]["revision"])

    from openpyxl import load_workbook
    book = load_workbook(fixture["attendance"])
    try:
        book["Arun"]["I4"] = 21
        book.save(fixture["attendance"])
    finally:
        book.close()

    result = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)
    assert result["protected"]
    assert repo.run(2026, 9)["state"] == PayrollState.CONFIRMED
    assert repo.db.query("SELECT * FROM source_changes WHERE run_id=?", (prepared["run"]["id"],))
    assert repo.audit_history(limit=1)[0]["action"] == "ATTENDANCE_CHANGED_AFTER_CONFIRMATION"


def test_phase_e_lifecycle_requires_reason_for_reopen_and_written_state_for_lock():
    fixture, _, service, _ = _service_with_configured_synthetic_inputs("lifecycle")
    prepared = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)
    service.confirm(2026, 9, actor="phase-e-uAT", expected_revision=prepared["run"]["revision"])
    with pytest.raises(PayrollStateError):
        service.reopen(2026, 9, reason="")
    with pytest.raises(PayrollStateError):
        service.lock(2026, 9)
    service.reopen(2026, 9, reason="Synthetic UAT review", actor="phase-e-uAT")
    assert service.current(2026, 9)[0]["state"] == PayrollState.DRAFT


def test_phase_e_same_synthetic_pair_is_idempotent_without_duplicate_audit_event():
    fixture, repo, service, _ = _service_with_configured_synthetic_inputs("idempotent")
    first = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)
    before = repo.db.query("SELECT COUNT(*) AS count FROM audit_log")[0]["count"]
    second = service.prepare(fixture["attendance"], fixture["payroll"], year=2026, month=9)
    after = repo.db.query("SELECT COUNT(*) AS count FROM audit_log")[0]["count"]
    assert second["unchanged"]
    assert second["run"]["revision"] == first["run"]["revision"]
    assert after == before


def test_phase_e_holiday_work_and_holiday_ot_need_explicit_evidence():
    fixture, repo, service, _ = _service_with_configured_synthetic_inputs("holiday_gate")
    from openpyxl import load_workbook
    book = load_workbook(fixture["attendance"])
    try:
        book["Arun"]["G9"] = 2  # 2026-09-06, a Sunday in this synthetic fixture
        book.save(fixture["attendance"])
    finally:
        book.close()
    arun = next(person for person in repo.employees() if person["canonical_name"] == "Arun Demo")
    repo.confirm_calendar("2026-09-06", "PUBLIC_HOLIDAY", "Synthetic holiday", required_workday=False,
                         actor="phase-e-uAT", source="synthetic-uAT")
    blocked = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)
    ot = next(issue for issue in blocked["draft"].issues if issue.kind == "OT")
    holiday_work = [issue for issue in blocked["draft"].issues if issue.kind == "HOLIDAY_WORK"]
    assert len(holiday_work) == 3
    repo.decide(blocked["draft"].source_hash, ot.key, {"category": "HOLIDAY_OT", "hours": "2"},
                reason="Synthetic holiday OT", actor="phase-e-uAT")
    for issue in holiday_work:
        repo.decide(blocked["draft"].source_hash, issue.key, {"holiday_days": "1"},
                    reason="Synthetic holiday work day", actor="phase-e-uAT")
    ready = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)["draft"]
    employee = next(employee for employee in ready.employees if employee.employee_id == arun["id"])
    assert ready.state == PayrollState.READY
    assert employee.values["holiday_ot"].final_value == 2
    assert employee.values["holiday_days"].final_value == 1


def test_phase_e_leave_decision_changes_bonus_only_when_explicitly_approved():
    fixture, repo, service, _ = _service_with_configured_synthetic_inputs("leave_gate")
    from openpyxl import load_workbook
    book = load_workbook(fixture["attendance"])
    try:
        book["Arun"]["J4"] = "ลา"
        book.save(fixture["attendance"])
    finally:
        book.close()
    blocked = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)
    leave = next(issue for issue in blocked["draft"].issues if issue.kind == "LEAVE")
    arun = next(employee for employee in blocked["draft"].employees if employee.name == "Arun Demo")
    assert arun.values["bonus"].final_value is None
    repo.decide(blocked["draft"].source_hash, leave.key, {"leave_type": "SICK_LEAVE", "bonus_effect": "CUT"},
                reason="Synthetic leave decision", actor="phase-e-uAT")
    ready = service.prepare(fixture["attendance"], fixture["payroll"], force=True, year=2026, month=9)["draft"]
    arun = next(employee for employee in ready.employees if employee.name == "Arun Demo")
    assert ready.state == PayrollState.READY
    assert arun.values["bonus"].final_value == 0
