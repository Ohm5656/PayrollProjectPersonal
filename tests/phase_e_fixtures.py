"""Synthetic-only Excel fixtures for controlled Phase E UAT.

These builders never read or copy a source workbook.  Callers must provide a
new child of ``test_workspace/uat_synthetic`` so that generated workbooks and
any later writer artefacts cannot overlap normal, QA, or historical fixtures.
"""
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook

from app.analysis.historical import FIELD_LABELS
from app.excel.attendance_reader import HEADERS
from app.safety import TEST_ROOT, SafetyError


PEOPLE = (
    ("Arun", "Arun Demo"),
    ("Bua", "Bua Example"),
    ("Chai", "Chai Fixture"),
)


def uat_root(label="case"):
    """Return a unique, non-existing safe UAT directory without deleting data."""
    root = (TEST_ROOT / "uat_synthetic" / f"{label}_{uuid4().hex}").resolve()
    if TEST_ROOT.resolve() not in root.parents:
        raise SafetyError("Synthetic UAT root escaped test_workspace")
    root.mkdir(parents=True)
    return root


def _header(field):
    return next(text for text, value in HEADERS.items() if value == field)


def _attendance(path: Path):
    book = Workbook()
    book.remove(book.active)
    for sheet_name, full_name in PEOPLE:
        sheet = book.create_sheet(sheet_name)
        sheet["A1"] = "ชื่อ"
        sheet["B1"] = full_name
        sheet["J1"] = datetime(2026, 9, 1)
        headers = ("date", "weekday", "work", "location", "clock_in", "clock_out", "ot", "allowance", "travel", "note")
        for column, field in enumerate(headers, 1):
            sheet.cell(3, column, _header(field))
        for day in range(1, 31):
            row = day + 3
            sheet.cell(row, 1, date(2026, 9, day))
            sheet.cell(row, 2, "วันทำงาน")
            sheet.cell(row, 3, "Synthetic office")
            sheet.cell(row, 4, "UAT site")
            sheet.cell(row, 5, time(8, 0))
            sheet.cell(row, 6, time(17, 0))
            sheet.cell(row, 7, Decimal("0"))
            sheet.cell(row, 8, Decimal("20"))
            sheet.cell(row, 9, Decimal("10"))
    book.save(path)
    book.close()


def _payroll(path: Path):
    book = Workbook()
    sheet = book.active
    sheet.title = "คิดค่าจ้าง"
    sheet["A1"] = "ประจำเดือน กันยายน 2569 จ่ายวันที่ 30 กันยายน 2569"
    sheet["A4"] = "ลำดับ"
    sheet["B4"] = "ชื่อ-สกุล"
    sheet["C4"] = "ค่าจ้าง"
    for start, title, labels in (
        (4, "ค่าล่วงเวลาปกติ x 1.5", ("ชั่วโมงละ", "จำนวนชม.", "รวมเงิน")),
        (7, "ค่าทำงานในวันหยุด", ("วันละ", "จำนวนวัน", "รวมเงิน")),
        (10, "OT วันหยุด x 3", ("ชั่วโมงละ", "จำนวนชม.", "รวมเงิน")),
    ):
        sheet.cell(4, start, title)
        for offset, label in enumerate(labels):
            sheet.cell(5, start + offset, label)
    field_columns = {
        "bonus": 13, "travel": 14, "allowance": 15, "phone": 16,
        "position": 17, "housing": 18, "income": 19, "net": 20,
    }
    for field, column in field_columns.items():
        sheet.cell(4, column, FIELD_LABELS[field] if field in FIELD_LABELS else {"income": "รวมเงินได้", "net": "รายได้สุทธิ"}[field])
    for ordinal, (_, full_name) in enumerate(PEOPLE, 1):
        row = ordinal + 6
        sheet.cell(row, 1, ordinal)
        sheet.cell(row, 2, full_name)
        values = {3: 30000, 4: 100, 5: 2, 7: 500, 8: 1, 10: 300, 11: 3,
                  13: 1000, 14: 250, 15: 350, 16: 500, 17: 700, 18: 900}
        for column, value in values.items():
            sheet.cell(row, column, value)
        sheet.cell(row, 6, f"=D{row}*E{row}")
        sheet.cell(row, 9, f"=G{row}*H{row}")
        sheet.cell(row, 12, f"=J{row}*K{row}")
        sheet.cell(row, 19, f"=C{row}+F{row}+I{row}+L{row}+M{row}+N{row}+O{row}+P{row}+Q{row}")
        sheet.cell(row, 20, f"=S{row}-R{row}")
    _payslip_sheets(book)
    book.save(path)
    book.close()


def _payslip_sheets(book):
    for _, full_name in PEOPLE:
        sheet = book.create_sheet(f"Slip {full_name.split()[0]}")
        sheet["A1"] = "PAYSLIP"
        sheet["A2"] = "Emp."
        sheet["B2"] = full_name
        sheet["A4"] = "รายได้ (Income)"
        sheet.merge_cells("A4:B4")
        sheet["D4"] = "รายการหัก (Deduction)"
        sheet.merge_cells("D4:E4")
        for row, label, value in (
            (5, "เงินเดือนค่าจ้าง", 30000), (6, "ค่าล่วงเวลา (100 x 2)", 200),
            (7, "ค่าทำงานวันหยุด (500 x 1)", 500), (8, "ค่าล่วงเวลาวันหยุด (300 x 3)", 900),
            (9, "เบี้ยขยัน", 1000), (10, "ค่าเดินทาง", 250), (11, "เบี้ยเลี้ยง", 350),
        ):
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, value)
        sheet["A12"] = "ค่าโทรศัพท์"
        sheet["B12"] = 500
        sheet["A13"] = "ค่าตำแหน่ง"
        sheet["B13"] = 700
        sheet["A14"] = "รายได้สะสม"
        sheet["D5"] = "ที่พัก+ค่าน้ำ+ไฟ"
        sheet["E5"] = 900
        sheet["D12"] = "ภาษีสะสม"
        sheet["A16"] = "รวมเงินได้"
        sheet["B16"] = 34400
        sheet["A17"] = "รวมเงินหัก"
        sheet["B17"] = 900
        sheet["A18"] = "เงินได้สุทธิ"
        sheet["B18"] = 33500
        sheet["A20"] = "period"
        sheet["B20"] = "September 2026"
        sheet["A21"] = "(date)"
        sheet["B21"] = "2026-09-30"


def create_synthetic_excel_fixture(root: Path | None = None):
    """Build only new synthetic input/output artefacts and return their paths."""
    root = Path(root or uat_root()).resolve()
    expected_parent = (TEST_ROOT / "uat_synthetic").resolve()
    if expected_parent not in root.parents:
        raise SafetyError("Phase E fixtures must live below test_workspace/uat_synthetic")
    root.mkdir(parents=True, exist_ok=False) if not root.exists() else None
    attendance_dir, payroll_dir = root / "attendance", root / "payroll_output"
    attendance_dir.mkdir(exist_ok=False)
    payroll_dir.mkdir(exist_ok=False)
    attendance = attendance_dir / "synthetic_attendance_9.xlsx"
    payroll = payroll_dir / "synthetic_payroll_9.xlsx"
    _attendance(attendance)
    _payroll(payroll)
    return {"root": root, "attendance": attendance, "payroll": payroll, "output": payroll_dir}
